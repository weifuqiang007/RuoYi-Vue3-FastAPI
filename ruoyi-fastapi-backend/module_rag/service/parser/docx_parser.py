# module_rag/service/parser/docx_parser.py
"""DOCX 文档解析器
参考 ragflow/deepdoc/parser/docx_parser.py 的 RAGFlowDocxParser 类
提取段落文字和表格内容
"""
from docx import Document


class DocxParser:

    def parse(self, file_path: str) -> list[dict]:
        """
        解析DOCX文件
        返回：[{"text":"内容", "page":0, "type":"text/title/table"}]
        """
        doc = Document(file_path)
        blocks = []
        paragraphs_by_element = {paragraph._element: paragraph for paragraph in doc.paragraphs}

        for element in doc.element.body:
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag
            if tag == "p":
                paragraph = paragraphs_by_element.get(element)
                para_text = paragraph.text if paragraph else ''
                if not para_text.strip():
                    continue

                # 判断是否标题（参考 ragflow 的 docx_question_level 函数）
                para_style = paragraph.style.name if paragraph and paragraph.style else ''

                is_heading = 'heading' in para_style.lower()
                block = {
                    "text": para_text.strip(),
                    "page": 0,
                    "type": 'title' if is_heading else 'text',
                }
                if is_heading:
                    try:
                        block["heading_level"] = int(para_style.split()[-1])
                    except (TypeError, ValueError):
                        block["heading_level"] = 1
                blocks.append(block)

            elif tag == 'tbl':
                # 表格处理 — 参考 ragflow 的 __extract_table_content 方法
                table_text = self._extract_table(element, doc)
                if table_text.strip():
                    blocks.append({
                        "text": table_text,
                        "page": 0,
                        "type": "table",
                    })

        return blocks

    def _extract_table(self, tbl_element, doc) -> str:
        """提取表格文本 — 参考 ragflow 的 __extract_table_content"""
        for table in doc.tables:
            if table._element is tbl_element:
                rows = []
                for row in table.rows:
                    rows.append(' | '.join(cell.text.strip() for cell in row.cells))
                return '\n'.join(rows)
        return ""
