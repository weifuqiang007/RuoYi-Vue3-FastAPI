from docx import Document

from module_rag.service.chunker.fixed_chunker import FixedChunker
from module_rag.service.parser.docx_parser import DocxParser
from module_rag.service.parser.markdown_parser import MarkdownParser


def test_markdown_parser_preserves_heading_path_and_block_type(tmp_path):
    markdown = tmp_path / 'sample.md'
    markdown.write_text(
        '# 第一章\n\n## 权力关系\n\n> 一段访谈引用\n\n- 价值澄清\n- 角色反思\n',
        encoding='utf-8',
    )

    blocks = MarkdownParser().parse(str(markdown))

    quote = next(block for block in blocks if block['type'] == 'quote')
    list_block = next(block for block in blocks if block['type'] == 'list')
    assert quote['heading_path'] == ['第一章', '权力关系']
    assert list_block['heading_path'] == ['第一章', '权力关系']
    assert blocks[0]['text'] == '第一章'


def test_docx_parser_reads_all_body_elements(tmp_path):
    path = tmp_path / 'sample.docx'
    document = Document()
    document.add_heading('章节标题', level=1)
    document.add_paragraph('第一段正文')
    document.add_paragraph('第二段正文')
    document.save(path)

    blocks = DocxParser().parse(str(path))

    assert [block['text'] for block in blocks] == ['章节标题', '第一段正文', '第二段正文']


def test_long_paragraph_chunking_does_not_duplicate_tail():
    text = '甲' * 260
    chunks = FixedChunker(chunk_size=100, overlap=10).chunk(text)

    assert len(chunks) == 3
    assert chunks[-1]['content'] != chunks[-2]['content'][-10:]
